from pathlib import Path

from .logger import logger

# These are placeholders that will be populated during lazy loading in load_models()
docx = None
np = None
pypdf = None
torch = None
Image = None
SentenceTransformer = None
AutoModel = None
AutoProcessor = None
AutoTokenizer = None
cosine_similarity = None


class MultimodalFileOrganizer:
    def __init__(self, categories_config=None):
        self.device = self._get_device()
        self.categories_config = categories_config or {}
        self.text_model = None
        self.image_model = None
        self.image_processor = None
        self.text_category_embeddings = {}

        # Flags
        self.models_loaded = False

    def _get_device(self):
        """Detects the best available device (CUDA, MPS, or CPU)."""
        try:
            if torch.cuda.is_available():
                return "cuda"
            elif torch.backends.mps.is_available():
                return "mps"
        except (NameError, AttributeError):
            pass
        return "cpu"

    def are_models_present(self):
        """
        Checks if both required models are likely present in the cache.
        """
        try:
            from transformers import AutoConfig
            # Check SigLIP 2
            AutoConfig.from_pretrained("google/siglip2-base-patch32-256", local_files_only=True)
            # Check Qwen Embedding
            AutoConfig.from_pretrained("Qwen/Qwen3-Embedding-0.6B", local_files_only=True)
            return True
        except Exception:
            return False

    def ensure_models(self, progress_callback=None):
        self.load_models(progress_callback)

    def load_models(self, progress_callback=None):
        """Import heavy libraries and load models from disk/HuggingFace."""
        global docx, np, pypdf, torch, Image, SentenceTransformer
        global AutoModel, AutoProcessor, AutoTokenizer, cosine_similarity
        if self.models_loaded:
            if progress_callback:
                progress_callback("Models already loaded.", 1.0)
            return

        # Lazy imports of heavy dependencies
        global docx, np, pypdf, torch, Image, SentenceTransformer, AutoModel, AutoProcessor, AutoTokenizer
        try:
            import docx as docx_mod
            import numpy as np_mod
            import pypdf as pypdf_mod
            import torch as torch_mod
            from PIL import Image as Image_mod
            from sentence_transformers import SentenceTransformer as SentenceTransformer_cls
            from transformers import AutoModel as AutoModel_cls
            from transformers import AutoProcessor as AutoProcessor_cls
            from transformers import AutoTokenizer as AutoTokenizer_cls

            docx = docx_mod
            np = np_mod
            pypdf = pypdf_mod
            torch = torch_mod
            Image = Image_mod
            SentenceTransformer = SentenceTransformer_cls
            AutoModel = AutoModel_cls
            AutoProcessor = AutoProcessor_cls
            AutoTokenizer = AutoTokenizer_cls

            from sklearn.metrics.pairwise import cosine_similarity as cosine_similarity_func
            cosine_similarity = cosine_similarity_func

            self.models_loaded = True

        except ImportError as e:
            logger.error(f"Failed to import ML dependencies: {e}")
            raise e

        try:
            if progress_callback:
                progress_callback("Loading Text Model (Qwen)...", 0.1)

            # Load Text Model
            self.text_model = SentenceTransformer(
                "Qwen/Qwen3-Embedding-0.6B",
                device=self.device,
                trust_remote_code=True
            )

            if progress_callback:
                progress_callback("Loading Image Model (SigLIP)...", 0.4)

            # Load Image Model
            self.image_model = AutoModel.from_pretrained(
                "google/siglip2-base-patch32-256",
            ).to(self.device).eval()

            self.image_processor = AutoProcessor.from_pretrained(
                "google/siglip2-base-patch32-256"
            )

            if progress_callback:
                progress_callback("Precomputing embeddings...", 0.8)

            self._precompute_text_embeddings()

            if progress_callback:
                progress_callback("Models loaded.", 1.0)

            self.models_loaded = True

        except Exception as e:
            logger.error(f"Error loading ML models: {e}")
            raise e

    def _precompute_text_embeddings(self):
        """Precompute category embeddings for text"""
        self.text_category_embeddings = {}
        instruction = "Instruct: Classify this content into file categories\nQuery:"

        if not self.categories_config:
            return

        for cat, desc in self.categories_config.items():
            if "text" in desc:
                text_to_encode = f"{instruction}{desc['text']}"
                emb = self.text_model.encode(
                    text_to_encode,
                    prompt_name="query",
                    convert_to_numpy=True
                )
                self.text_category_embeddings[cat] = emb

    def extract_text(self, file_path: Path):
        """Extracts text from various file formats."""
        ext = file_path.suffix.lower()
        content = ""

        try:
            if ext in ['.txt', '.md', '.py', '.js', '.html', '.css', '.json', '.xml']:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read(5000) # Limit to first 5KB

            elif ext == '.pdf':
                try:
                    reader = pypdf.PdfReader(file_path)
                    # Extract text from first few pages
                    for i in range(min(3, len(reader.pages))):
                        page_text = reader.pages[i].extract_text()
                        if page_text:
                            content += page_text + "\n"
                except Exception as e:
                    logger.error(f"PDF extraction error: {e}")

            elif ext == '.docx':
                try:
                    doc = docx.Document(file_path)
                    # Limit paragraphs
                    paragraphs = [p.text for p in doc.paragraphs[:50]]
                    content = "\n".join(paragraphs)
                except Exception as e:
                    logger.error(f"Docx extraction error: {e}")

        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {e}")

        return content

    def categorize_image(self, image_path):
        """Categorize image using SigLIP 2"""
        if not self.models_loaded:
            return None, 0.0

        try:
            image = Image.open(image_path).convert("RGB")

            # Collect all visual descriptions
            all_labels = []
            label_to_category = {}

            for cat, desc in self.categories_config.items():
                if 'visual' in desc:
                    visual_descs = desc['visual']
                    if isinstance(visual_descs, str):
                        visual_descs = [visual_descs]

                    for label in visual_descs:
                        all_labels.append(label)
                        label_to_category[label] = cat

            if not all_labels:
                return None, 0.0

            # Prepare inputs
            inputs = self.image_processor(
                images=image,
                text=all_labels,
                return_tensors="pt",
                padding="max_length"
            ).to(self.device)

            # Get predictions
            with torch.no_grad():
                outputs = self.image_model(**inputs)
                probs = torch.sigmoid(outputs.logits_per_image[0])

            # Find best match
            best_idx = probs.argmax().item()
            best_label = all_labels[best_idx]
            confidence = probs[best_idx].item()
            category = label_to_category[best_label]

            return category, confidence

        except Exception as e:
            logger.error(f"Error categorizing image {image_path}: {e}")
            return None, 0.0

    def categorize_text_file(self, file_path, content, threshold=0.4):
        """Categorize text-based file using Qwen3"""
        global cosine_similarity
        if not self.models_loaded or not content or len(content.strip()) < 10:
            return None, 0.0

        try:
            instruction = "Instruct: Classify this content into file categories\nQuery:"
            # Qwen embedding
            content_emb = self.text_model.encode(
                f"{instruction}{content[:2000]}",
                prompt_name="query",
                convert_to_numpy=True
            )

            if cosine_similarity is None:
                from sklearn.metrics.pairwise import cosine_similarity as cosine_similarity_func
                cosine_similarity = cosine_similarity_func

            # Compute similarities
            similarities = {}
            for cat, cat_emb in self.text_category_embeddings.items():
                # Cosine similarity
                sim = np.dot(content_emb, cat_emb) / (
                    np.linalg.norm(content_emb) * np.linalg.norm(cat_emb) + 1e-9
                )
                similarities[cat] = sim

            if not similarities:
                return None, 0.0

            best_category = max(similarities.items(), key=lambda x: x[1])
            return best_category[0], float(best_category[1])

        except Exception as e:
            logger.error(f"Error categorizing text {file_path}: {e}")
            return None, 0.0

    def smart_categorize(self, file_path, threshold=0.3):
        """Main categorization method"""
        if not self.models_loaded:
            return None, 0.0, "ml-not-loaded"

        file_ext = file_path.suffix.lower()

        # Image files
        if file_ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.webp']:
            category, confidence = self.categorize_image(file_path)
            if confidence > threshold:
                return category, confidence, "image-ml"

        # Text-extractable files
        elif file_ext in ['.txt', '.md', '.py', '.js', '.html', '.pdf', '.docx', '.css', '.json']:
            content = self.extract_text(file_path)
            if content:
                category, confidence = self.categorize_text_file(file_path, content)
                if confidence > threshold:
                    return category, confidence, "text-ml"

        return None, 0.0, "extension"
